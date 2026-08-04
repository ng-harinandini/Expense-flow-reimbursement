"""Microsoft Word (``.docx``) via ``python-docx``.

Word has no native page concept — layout is reflowed at render time, not stored in the file — so
``DocumentSection.page_number`` stays ``None`` here; only the PDF and image-OCR parsers can
populate it. Structure instead comes from paragraph styles: ``Heading 1``…``Heading 9`` (and
``Title``, folded to level 1) drive the same ancestor-tracking helper
(:func:`~app.ai.parsing.base.sections_from_headings`) the HTML and Markdown parsers use, so a Word
policy document cites exactly as precisely as an HTML one.

Tables are appended as plain rows (tab-joined cells, one row per line) into whichever section they
appear in, rather than dropped — a per-country per-diem rate table is exactly the kind of content
an expense policy needs to be retrievable, and a table has no heading structure of its own to
preserve.

**Body content is walked in true document order, not paragraphs-then-tables.**
``Document.paragraphs`` and ``Document.tables`` are two separate flat lists python-docx exposes for
convenience, and neither carries the other's interleaving — a document with headings A, a table,
then heading B loses that ordering if the two lists are consumed one after the other, and every
table in the document ends up attributed to whichever heading happened to be current when the
*second* loop ran (in practice, the last one). Walking ``document.element.body`` directly and
dispatching on each child's own tag is what preserves the real order.
"""

from __future__ import annotations

import io
import re
from typing import Sequence

from app.ai.core.types import DocumentSection, RawDocument
from app.ai.parsing.base import DEFAULT_MAX_BYTES, BaseDocumentParser, sections_from_headings

# python-docx's own body-element tags: <w:p> is a paragraph, <w:tbl> is a table. Compared against
# ``element.tag`` directly (a namespaced string python-docx computes via ``qn()`` internally) rather
# than importing ``docx.oxml.ns.qn`` at module scope, which would defeat the guarded-import contract
# below by making importing *this module* fail outright on a host without python-docx installed.
_PARAGRAPH_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
_TABLE_TAG = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl"

NAME = "docx"
MIME_TYPES = frozenset({
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
})
EXTENSIONS = frozenset({".docx"})

_HEADING_STYLE = re.compile(r"^Heading (\d+)$")


class DocxParser(BaseDocumentParser):
    """Word documents, sectioned by paragraph heading style."""

    def __init__(self, *, max_bytes: int = DEFAULT_MAX_BYTES, recorder=None) -> None:
        super().__init__(
            name=NAME, mime_types=MIME_TYPES, extensions=EXTENSIONS,
            max_bytes=max_bytes, recorder=recorder,
        )

    def is_available(self) -> bool:
        try:
            import docx  # noqa: F401, PLC0415
        except ImportError:
            return False
        return True

    def _parse(self, document: RawDocument) -> tuple[str, Sequence[DocumentSection]]:
        from docx import Document

        doc = Document(io.BytesIO(document.content))
        entries = list(_entries(doc))
        sections = sections_from_headings(entries)
        text = "\n\n".join(s.text for s in sections)
        return text, sections


def _entries(doc) -> Sequence[tuple[int, str, str]]:  # noqa: ANN001
    """Walk the document body once, in the order Word itself stores it.

    ``doc.element.body`` yields raw XML elements; each is wrapped back into the ``Paragraph``/
    ``Table`` object python-docx would have given us from ``doc.paragraphs``/``doc.tables``, so a
    table between two headings lands in the section it actually appears under rather than being
    attributed to whichever heading happened to be current once every paragraph had already been
    consumed.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    entries: list[tuple[int, str, str]] = []
    current_level, current_heading, current_lines = 0, "", []

    for child in doc.element.body:
        if child.tag == _PARAGRAPH_TAG:
            paragraph = Paragraph(child, doc)
            level = _heading_level(paragraph.style.name if paragraph.style else "")
            text = paragraph.text.strip()
            if level is not None:
                entries.append((current_level, current_heading, "\n".join(current_lines).strip()))
                current_level = level
                current_heading = text
                current_lines = []
            elif text:
                current_lines.append(text)
        elif child.tag == _TABLE_TAG:
            current_lines.append(_table_text(Table(child, doc)))

    entries.append((current_level, current_heading, "\n".join(current_lines).strip()))
    return [e for e in entries if e[1] or e[2]] or [(0, "", "")]


def _heading_level(style_name: str) -> "int | None":
    if style_name == "Title":
        return 1
    match = _HEADING_STYLE.match(style_name or "")
    return int(match.group(1)) if match else None


def _table_text(table) -> str:  # noqa: ANN001
    rows = [
        "\t".join(cell.text.strip() for cell in row.cells)
        for row in table.rows
    ]
    return "\n".join(row for row in rows if row.strip())


__all__ = ["EXTENSIONS", "MIME_TYPES", "NAME", "DocxParser"]
